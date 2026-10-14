"""
build_docx.py - builds the EDITABLE Word version of the report (IVP_Mini_Project_Report.docx)
from the same text as the PDF (report/content.py) and the same figures (report/figs).

Run from the project folder:   venv\\Scripts\\python.exe report\\build_docx.py
"""

import html
import json
import os
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
FIG = os.path.join(HERE, "figs")
OUT = os.path.join(HERE, "IVP_Mini_Project_Report.docx")
D = json.load(open(os.path.join(FIG, "data.json")))

# title-page details: same placeholders as the PDF (edit them directly in Word)
DETAILS = {
    "students": [("[Student Name 1]", "[Roll No. 1]"),
                 ("[Student Name 2]", "[Roll No. 2]"),
                 ("[Student Name 3]", "[Roll No. 3]")],
    "programme": "[Programme] / Semester [ ] / Division [ ]",
    "faculty": "[Faculty Name]",
    "institute": "SVKM's NMIMS Mukesh Patel School of Technology Management & Engineering (MPSTME), "
                 "[Department]",
    "year": "2026-27",
}

NAVY = RGBColor(0x1F, 0x3B, 0x63)
doc = Document()

# ---------------------------------------------------------------- page + styles
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)          # A4
sec.left_margin = sec.right_margin = Cm(2.0)
sec.top_margin, sec.bottom_margin = Cm(1.8), Cm(2.0)

normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.15
for name, size in (("Heading 1", 16), ("Heading 2", 13)):
    st = doc.styles[name]
    st.font.name = "Calibri"
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = NAVY

fig_no = [0]
tab_no = [0]


# ---------------------------------------------------------------- inline markup
TOKEN = re.compile(r"(<br\s*/?>|</?b>|</?i>|</?sub>|</?super>|<font[^>]*>|</font>)")


def add_rich(par, text, size=None, font=None, bold=False):
    """Add text with the simple ReportLab-style tags used in content.py."""
    state = {"b": bold, "i": False, "sub": False, "sup": False, "mono": False}
    for part in TOKEN.split(text):
        if not part:
            continue
        low = part.lower()
        if low.startswith("<br"):
            par.add_run().add_break()
        elif low in ("<b>", "</b>"):
            state["b"] = low == "<b>" or bold
        elif low in ("<i>", "</i>"):
            state["i"] = low == "<i>"
        elif low in ("<sub>", "</sub>"):
            state["sub"] = low == "<sub>"
        elif low in ("<super>", "</super>"):
            state["sup"] = low == "<super>"
        elif low.startswith("<font"):
            state["mono"] = "courier" in low
        elif low == "</font>":
            state["mono"] = False
        else:
            run = par.add_run(html.unescape(part).replace(" ", " "))
            run.bold = state["b"]
            run.italic = state["i"]
            run.font.subscript = state["sub"]
            run.font.superscript = state["sup"]
            if state["mono"] or font:
                run.font.name = "Consolas" if state["mono"] else font
            if size:
                run.font.size = Pt(size)
    return par


def shade(cell_or_par, hex_fill):
    el = cell_or_par._tc.get_or_add_tcPr() if hasattr(cell_or_par, "_tc") else cell_or_par._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    el.append(shd)


# ---------------------------------------------------------------- helpers used by content.py
def title_page():
    def line(text, size=12, bold=False, color=None, space=6):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(space)
        add_rich(p, text, size=size, bold=bold)
        if color:
            for r in p.runs:
                r.font.color.rgb = color
        return p

    for _ in range(3):
        doc.add_paragraph()
    line("Gesture IVP Lab", 26, True, NAVY, 4)
    line("A Hand-Gesture-Controlled Image Processing Laboratory", 14, space=6)
    line("Image and Video Processing - Mini-Project Report", 12, space=30)
    # the fields required by the report guideline, in its order (click a cell in Word to edit)
    names = DETAILS["students"]
    fields = [("Project Title", "Gesture IVP Lab: A Hand-Gesture-Controlled Image Processing Laboratory"),
              ("Course", "Image and Video Processing"),
              ("Student Name(s)", "<br/>".join(n for n, _ in names)),
              ("Roll Number(s)", "<br/>".join(r for _, r in names)),
              ("Programme / Semester / Division", DETAILS["programme"]),
              ("Faculty Name", DETAILS["faculty"]),
              ("Institute / Department", DETAILS["institute"]),
              ("Academic Year", DETAILS["year"])]
    t = doc.add_table(rows=len(fields), cols=2)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r, (k, v) in enumerate(fields):
        a, b = t.rows[r].cells
        a.width, b.width = Cm(5.5), Cm(11)
        shade(a, "DFE8F3")
        for cell, txt in ((a, f"<b>{k}</b>"), (b, v)):
            cell.paragraphs[0].paragraph_format.space_before = Pt(4)
            cell.paragraphs[0].paragraph_format.space_after = Pt(4)
            add_rich(cell.paragraphs[0], txt, size=12)
    pagebreak()


def h1(text):
    doc.add_heading(html.unescape(re.sub(r"<[^>]+>", "", text)), level=1)


def h2(text):
    doc.add_heading(html.unescape(re.sub(r"<[^>]+>", "", text)), level=2)


def P(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    add_rich(p, text)


def bullets(items):
    for it in items:
        add_rich(doc.add_paragraph(style="List Bullet"), it)


def numbered(items):
    """Numbered list 1., 2., 3. ... (typed numbers, so every list restarts at 1)."""
    for i, it in enumerate(items, 1):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.left_indent = Cm(0.75)
        p.paragraph_format.first_line_indent = Cm(-0.6)
        p.paragraph_format.space_after = Pt(3)
        add_rich(p, f"{i}.  {it}")


def figure(fname, caption, width=17):
    fig_no[0] += 1
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(os.path.join(FIG, fname), width=Cm(width))
    c = doc.add_paragraph()
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_rich(c, f"<b>Figure {fig_no[0]}.</b> {caption}", size=9.5)


def table(rows, widths, caption, header=True):
    tab_no[0] += 1
    c = doc.add_paragraph()
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c.paragraph_format.keep_with_next = True
    add_rich(c, f"<b>Table {tab_no[0]}.</b> {caption}", size=9.5)
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = t.rows[r].cells[ci]
            cell.width = Cm(widths[ci])
            para = cell.paragraphs[0]
            para.paragraph_format.space_after = Pt(0)
            para.paragraph_format.keep_with_next = r < len(rows) - 1   # keep the table on one page
            add_rich(para, str(val), size=9.5, bold=(header and r == 0))
            if header and r == 0:
                shade(cell, "DFE8F3")
    doc.add_paragraph()


def code(text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.4)
    shade(p, "F2F2F2")
    add_rich(p, text, size=9, font="Consolas")


def pagebreak():
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def ref(text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.8)
    p.paragraph_format.first_line_indent = Cm(-0.8)
    add_rich(p, text)


# ---------------------------------------------------------------- footer with page number
def add_page_number(section):
    p = section.footer.paragraphs[0]
    p.text = "Gesture IVP Lab - Image and Video Processing mini-project\tPage "
    for r in p.runs:
        r.font.size = Pt(8.5)
    run = p.add_run()
    run.font.size = Pt(8.5)
    for kind, txt in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = txt
        run._r.append(el)
    section.different_first_page_header_footer = True     # no footer on the title page


# ---------------------------------------------------------------- build
ns = dict(title_page=title_page, h1=h1, h2=h2, P=P, bullets=bullets, numbered=numbered, figure=figure, table=table,
          code=code, pagebreak=pagebreak, ref=ref, D=D, DETAILS=DETAILS)
exec(compile(open(os.path.join(HERE, "content.py"), encoding="utf-8").read(), "content.py", "exec"), ns)
add_page_number(sec)
doc.core_properties.title = "Gesture IVP Lab - IVP Mini-Project Report"
doc.save(OUT)
print("saved", OUT)
